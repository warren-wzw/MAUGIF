import os
import sys
os.chdir(sys.path[0])
import numpy as np
from PIL import Image
from Metric import *
from natsort import natsorted
from tqdm import tqdm
import warnings
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")
DATASET="MSRS"
METHOD = ["EMDiFusion"]
INFRARED=f"../source_image/{DATASET}/ir"
VISIBLE=f"../source_image/{DATASET}/vi"
ROUND=6

def write_excel(excel_name='metric.xlsx', worksheet_name='VIF', column_index=0, data=None):
    try:
        workbook = load_workbook(excel_name)
    except FileNotFoundError:
        workbook = Workbook()
    if worksheet_name in workbook.sheetnames:
        worksheet = workbook[worksheet_name]
    else:
        worksheet = workbook.create_sheet(title=worksheet_name)
    column = get_column_letter(column_index + 1)
    for i, value in enumerate(data):
        cell = worksheet[column + str(i + 1)]
        cell.value = value
    workbook.save(excel_name)

def evaluation_one(ir_name, vi_name, f_name):
    f_img = Image.open(f_name).convert('L')
    ir_img = Image.open(ir_name).convert('L')
    vi_img = Image.open(vi_name).convert('L')
    f_img_int = np.array(f_img).astype(np.int32)
    f_img_double = np.array(f_img).astype(np.float32)
    ir_img_int = np.array(ir_img).astype(np.int32)
    ir_img_double = np.array(ir_img).astype(np.float32)
    vi_img_int = np.array(vi_img).astype(np.int32)
    vi_img_double = np.array(vi_img).astype(np.float32)

    EN = EN_function(f_img_int)
    MI = MI_function(ir_img_int, vi_img_int, f_img_int, gray_level=256)
    SF = SF_function(f_img_double)
    SD = SD_function(f_img_double)
    AG = AG_function(f_img_double)
    PSNR = PSNR_function(ir_img_double, vi_img_double, f_img_double)
    MSE = MSE_function(ir_img_double, vi_img_double, f_img_double)
    VIF = VIF_function(ir_img_double, vi_img_double, f_img_double)
    CC = CC_function(ir_img_double, vi_img_double, f_img_double)
    SCD = SCD_function(ir_img_double, vi_img_double, f_img_double)
    Qabf = Qabf_function(ir_img_double, vi_img_double, f_img_double)
    Nabf = Nabf_function(ir_img_double, vi_img_double, f_img_double)
    SSIM = SSIM_function(ir_img_double, vi_img_double, f_img_double)
    MS_SSIM = MS_SSIM_function(ir_img_double, vi_img_double, f_img_double)
    return EN, MI, SF, AG, SD, CC, SCD, VIF, MSE, PSNR, Qabf, Nabf, SSIM, MS_SSIM


if __name__ == '__main__':
    
    ir_dir = INFRARED
    vi_dir = VISIBLE
    for method in METHOD:
        Method = method
        f_dir = os.path.join(f'../Results/{DATASET}/', Method)
        save_dir = f'../Excel/{DATASET}'
        os.makedirs(save_dir, exist_ok=True)
        EN_list, MI_list, SF_list, AG_list, SD_list,CC_list, SCD_list, VIF_list,MSE_list,PSNR_list,Qabf_list,\
        Nabf_list,SSIM_list, MS_SSIM_list, filename_list= [],[],[],[],[],[],[],[],[],[],[],[],[],[],['']
        names=["EN","MI","SF","AG","SD","CC","SCD","VIF","MSE","PSNR","Qabf","Nabf","SSIM","MS_SSIM"]
        name_lists = [EN_list, MI_list, SF_list, AG_list, SD_list, CC_list, SCD_list, VIF_list, MSE_list,     
                        PSNR_list, Qabf_list, Nabf_list, SSIM_list, MS_SSIM_list]
        metric_save_name = os.path.join(save_dir, 'metric_{}.xlsx'.format( Method))
        filelist = natsorted(os.listdir(ir_dir))
        eval_bar = tqdm(filelist)
        for _, item in enumerate(eval_bar):
            ir_name = os.path.join(ir_dir, item)
            vi_name = os.path.join(vi_dir, item)
            f_name = os.path.join(f_dir, item)
            EN, MI, SF, AG, SD, CC, SCD, VIF, MSE, PSNR, Qabf, Nabf, SSIM, MS_SSIM = evaluation_one(ir_name, vi_name,f_name)
            EN_list.append(EN)
            MI_list.append(MI)
            SF_list.append(SF)
            AG_list.append(AG)
            SD_list.append(SD)
            CC_list.append(CC)
            SCD_list.append(SCD)
            VIF_list.append(VIF)
            MSE_list.append(MSE)
            PSNR_list.append(PSNR)
            Qabf_list.append(Qabf)
            Nabf_list.append(Nabf)
            SSIM_list.append(SSIM)
            MS_SSIM_list.append(MS_SSIM)
            filename_list.append(item)
            eval_bar.set_description("{} | {}".format(Method, item))
        filename_list.insert(1,"mean")
        for name,name_list in zip(names,name_lists):
            name_list.insert(0, round(np.mean(name_list), ROUND)) 
            name_list = [round(x, ROUND) for x in name_list]
            name_list.insert(0, '{}'.format(Method)) 
            write_excel(metric_save_name, name, 0, filename_list)
            write_excel(metric_save_name, name, 1, name_list)
